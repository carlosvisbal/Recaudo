from io import BytesIO

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill, Border, Side
from openpyxl.utils import get_column_letter

from .money import format_cop


def exportar_consultas(recaudos, totales) -> bytes:
    wb = Workbook()
    thin = Border(
        left=Side(style="thin", color="D9CBB3"),
        right=Side(style="thin", color="D9CBB3"),
        top=Side(style="thin", color="D9CBB3"),
        bottom=Side(style="thin", color="D9CBB3"),
    )
    header_fill = PatternFill("solid", fgColor="0F4C4C")
    header_font = Font(color="FFFFFF", bold=True)
    money_font = Font(name="Calibri")

    ws = wb.active
    ws.title = "Recaudos"
    headers = [
        "Recibo",
        "Fecha",
        "Estado",
        "Actividad",
        "Padrino",
        "Familia / aportante",
        "Valor COP",
        "Medio de pago",
        "Concepto",
        "Referencia",
        "Registrado por",
        "Soportes",
    ]
    ws.append(headers)
    for cell in ws[1]:
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal="center")

    for recaudo in recaudos:
        ws.append(
            [
                recaudo.numero_formateado,
                recaudo.fecha.isoformat(),
                recaudo.get_estado_display(),
                recaudo.actividad_nombre,
                recaudo.padrino_nombre,
                recaudo.aportante_nombre,
                float(recaudo.valor),
                recaudo.get_medio_pago_display(),
                recaudo.concepto,
                recaudo.referencia_consignacion,
                recaudo.usuario_nombre,
                recaudo.soportes.count(),
            ]
        )

    for row in ws.iter_rows(min_row=2, max_col=12):
        for cell in row:
            cell.border = thin
            cell.font = money_font
        row[6].number_format = '"$"#,##0.00'

    widths = [12, 14, 14, 28, 28, 32, 16, 16, 32, 18, 22, 12]
    for i, width in enumerate(widths, start=1):
        ws.column_dimensions[get_column_letter(i)].width = width
    ws.auto_filter.ref = f"A1:L{ws.max_row}"
    ws.freeze_panes = "A2"

    def _totales_sheet(title, rows, key_name):
        sheet = wb.create_sheet(title)
        sheet.append([key_name, "Cantidad", "Total COP"])
        for cell in sheet[1]:
            cell.fill = header_fill
            cell.font = header_font
        for item in rows:
            sheet.append([item["nombre"], item["cantidad"], float(str(item["total"]))])
        for row in sheet.iter_rows(min_row=2, max_col=3):
            row[2].number_format = '"$"#,##0.00'
        sheet.column_dimensions["A"].width = 36
        sheet.column_dimensions["B"].width = 14
        sheet.column_dimensions["C"].width = 18

    _totales_sheet("Por actividad", totales.get("por_actividad") or [], "Actividad")
    _totales_sheet("Por padrino", totales.get("por_padrino") or [], "Padrino")
    _totales_sheet("Por familia", totales.get("por_familia") or [], "Familia")
    _totales_sheet("Por fecha", totales.get("por_fecha") or [], "Fecha")

    cruzado = wb.create_sheet("Padrino x actividad")
    cruzado.append(
        [
            "Padrino",
            "Actividad",
            "Pagos",
            "Total COP",
            "Familias asignadas",
            "Familias que pagaron",
            "Familias pendientes",
        ]
    )
    for cell in cruzado[1]:
        cell.fill = header_fill
        cell.font = header_font
    for item in totales.get("por_padrino_actividad") or []:
        cruzado.append(
            [
                item.get("padrino_nombre"),
                item.get("actividad_nombre"),
                item.get("cantidad") or 0,
                float(str(item.get("total") or 0)),
                item.get("familias_asignadas") or 0,
                item.get("familias_pagaron") or 0,
                item.get("familias_pendientes") or 0,
            ]
        )
    for row in cruzado.iter_rows(min_row=2, max_col=7):
        row[3].number_format = '"$"#,##0.00'
    cruzado.column_dimensions["A"].width = 32
    cruzado.column_dimensions["B"].width = 32
    cruzado.column_dimensions["C"].width = 12
    cruzado.column_dimensions["D"].width = 16
    cruzado.column_dimensions["E"].width = 20
    cruzado.column_dimensions["F"].width = 22
    cruzado.column_dimensions["G"].width = 20

    pendientes_sheet = wb.create_sheet("Pendientes")
    pendientes_sheet.append(["Padrino", "Actividad", "Familia", "Celular"])
    for cell in pendientes_sheet[1]:
        cell.fill = header_fill
        cell.font = header_font
    for item in (totales.get("pendientes") or {}).get("results") or []:
        pendientes_sheet.append(
            [
                item.get("padrino_nombre"),
                item.get("actividad_nombre"),
                item.get("familia_nombre"),
                item.get("telefono") or "",
            ]
        )
    pendientes_sheet.column_dimensions["A"].width = 32
    pendientes_sheet.column_dimensions["B"].width = 32
    pendientes_sheet.column_dimensions["C"].width = 32
    pendientes_sheet.column_dimensions["D"].width = 18

    resumen = wb.create_sheet("Resumen")
    resumen.append(["Indicador", "Valor"])
    resumen["A1"].fill = header_fill
    resumen["B1"].fill = header_fill
    resumen["A1"].font = header_font
    resumen["B1"].font = header_font
    resumen.append(["Total confirmado", format_cop(totales["total_confirmado"])])
    resumen.append(["Cantidad confirmados", totales["cantidad_confirmados"]])
    resumen.append(["Cantidad anulados (excluidos del total)", totales["cantidad_anulados"]])
    pendientes = totales.get("pendientes") or {}
    resumen.append(["Familias asignadas (filtro)", pendientes.get("familias_asignadas", "")])
    resumen.append(["Familias al día", pendientes.get("familias_al_dia", "")])
    resumen.append(["Pendientes de pago (familia × actividad)", pendientes.get("count", "")])
    resumen.column_dimensions["A"].width = 42
    resumen.column_dimensions["B"].width = 24

    output = BytesIO()
    wb.save(output)
    return output.getvalue()
